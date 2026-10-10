#!/usr/bin/env python3
"""
parser.py — Citation extraction from .docx and .tex/.bib documents.

Extracts every citation along with its surrounding context (the sentence
and neighboring sentences where it appears) for downstream LLM verification.
"""

import re
from pathlib import Path
from typing import Dict, List, Tuple


# ---------------------------------------------------------------------------
# .docx parsing
# ---------------------------------------------------------------------------

def _parse_docx_references(doc) -> Tuple[Dict[str, str], int]:
    """Locate the References section by detecting a sequential (1, 2, 3…) pattern.

    Returns (references_dict, paragraph_index where references begin).
    """
    ref_pattern = re.compile(r"^\s*(?:\[)?(\d+)(?:\])?[\.\s]\s+(.*)", re.DOTALL)
    paragraphs = doc.paragraphs
    total = len(paragraphs)
    references: Dict[str, str] = {}
    start_index = -1

    # STEP 1: find where "1. …" is followed by "2. …" within 10 paragraphs
    for i in range(total):
        text = paragraphs[i].text.strip()
        if not text:
            continue
        match = ref_pattern.match(text)
        if match and int(match.group(1)) == 1:
            for j in range(i + 1, min(i + 10, total)):
                next_text = paragraphs[j].text.strip()
                next_match = ref_pattern.match(next_text)
                if next_match and int(next_match.group(1)) == 2:
                    start_index = i
                    break
            if start_index != -1:
                break

    if start_index == -1:
        return {}, -1

    # STEP 2: parse from start_index onwards, strict sequential numbering
    current_id = 0
    for i in range(start_index, total):
        text = paragraphs[i].text.strip()
        if not text:
            continue
        match = ref_pattern.match(text)
        if match:
            found_id = int(match.group(1))
            if found_id == current_id + 1:
                current_id = found_id
                references[str(current_id)] = match.group(2)
            else:
                # non-sequential number → append to previous ref
                if str(current_id) in references:
                    references[str(current_id)] += " " + text
        else:
            if current_id > 0 and str(current_id) in references:
                references[str(current_id)] += " " + text

    return references, start_index


def _extract_docx_citations(doc, ref_db: Dict[str, str], stop_index: int) -> Tuple[str, List[dict]]:
    """Scan document body for (1), [1], (1, 2), [1-3] style citations."""
    cite_pattern = re.compile(r"(?:[\(\[])([\d\s,\-\–]+)(?:[\)\]])")
    full_text_parts: List[str] = []
    citations_grouped: Dict[str, dict] = {}

    for i, para in enumerate(doc.paragraphs):
        if stop_index != -1 and i >= stop_index:
            break
        text = para.text.strip()
        if not text:
            continue
        full_text_parts.append(text)

        for match in cite_pattern.finditer(text):
            content = match.group(1)
            # Skip years (4-digit numbers not in ref_db)
            if len(content) == 4 and content.isdigit() and content not in ref_db:
                continue

            citation_ids: List[str] = []
            for part in content.split(","):
                part = part.strip()
                if "-" in part or "\u2013" in part:
                    range_parts = re.split(r"[-\–]", part)
                    try:
                        start, end = int(range_parts[0]), int(range_parts[1])
                        if (end - start) < 50:
                            citation_ids.extend(str(k) for k in range(start, end + 1))
                    except (ValueError, IndexError):
                        pass
                elif part.isdigit():
                    citation_ids.append(part)

            valid_ids = [cid for cid in citation_ids if cid in ref_db]
            for cid in valid_ids:
                if cid not in citations_grouped:
                    citations_grouped[cid] = {
                        "key": cid,
                        "ref_text": ref_db[cid],
                        "usages": [],
                    }
                citations_grouped[cid]["usages"].append(text)

    # Build context blocks
    citations = []
    for key, data in citations_grouped.items():
        unique_contexts = list(dict.fromkeys(data["usages"]))  # dedupe preserving order
        citations.append({
            "key": key,
            "ref_text": data["ref_text"],
            "usage_count": len(data["usages"]),
            "unique_context_count": len(unique_contexts),
            "context_block": "\n---\n".join(unique_contexts),
        })

    return "\n".join(full_text_parts), citations


def extract_from_docx(docx_path: str) -> List[dict]:
    """Extract citations from a Word (.docx) document.

    Returns a list of citation dicts, each containing:
        key: citation identifier (e.g. "1", "2")
        ref_text: reference string from bibliography
        usage_count: how many times cited in body
        unique_context_count: unique context blocks
        context_block: combined context text for verification
    """
    try:
        import docx
    except ImportError:
        raise ImportError(
            "python-docx is required for .docx parsing. "
            "Install it: pip install python-docx"
        )

    doc = docx.Document(docx_path)
    ref_db, start_index = _parse_docx_references(doc)
    if not ref_db:
        return []

    full_text, citations = _extract_docx_citations(doc, ref_db, start_index)
    return citations


# ---------------------------------------------------------------------------
# .tex / .bib parsing
# ---------------------------------------------------------------------------

def _parse_bib(bib_path: str) -> Dict[str, dict]:
    """Parse a .bib file into a dict keyed by citation ID."""
    try:
        import bibtexparser
    except ImportError:
        raise ImportError(
            "bibtexparser is required for .bib parsing. "
            "Install it: pip install bibtexparser"
        )

    with open(bib_path, "r", encoding="utf-8") as f:
        bib_db = bibtexparser.load(f)
    return {entry["ID"]: entry for entry in bib_db.entries}


def _format_bib_entry(entry: dict) -> str:
    """Format a BibTeX entry dict into a human-readable string."""
    priority = [
        "author", "title", "journal", "booktitle", "year",
        "volume", "number", "pages", "publisher", "doi", "url",
    ]
    lines = [f"@{entry.get('ENTRYTYPE', 'unknown')}{{{entry.get('ID', 'unknown')}}}"]
    seen = set()
    for field in priority:
        if field in entry:
            val = str(entry[field]).replace("{", "").replace("}", "")
            lines.append(f"  {field.capitalize():12}: {val}")
            seen.add(field)
    for field, value in entry.items():
        if field in seen or field in ("ENTRYTYPE", "ID"):
            continue
        lines.append(f"  {field.capitalize():12}: {str(value).replace('{', '').replace('}', '')}")
    return "\n".join(lines)


def _find_section_heading(content: str, position: int) -> str:
    """Find the most recent LaTeX section heading before *position*."""
    pattern = re.compile(
        r"\\(chapter|section|subsection|subsubsection|paragraph)\*?\s*"
        r"(?:\[.*?\])?\s*"
        r"\{([^}]*(?:\{[^}]*\}[^}]*)*)\}",
        re.IGNORECASE,
    )
    matches = list(pattern.finditer(content[:position]))
    if not matches:
        return "[SECTION] Unknown"
    last = matches[-1]
    title = last.group(2)
    # Clean LaTeX commands
    title = re.sub(r"\\texorpdfstring\s*\{([^}]*)\}\s*\{[^}]*\}", r"\1", title)
    title = re.sub(r"\\[a-zA-Z]+\{([^}]*)\}", r"\1", title)
    title = re.sub(r"\\[a-zA-Z]+", "", title)
    title = re.sub(r"[{}]", "", title)
    title = re.sub(r"\s+", " ", title).strip()
    return f"[{last.group(1).upper()}] {title}"


def _split_sentences(text: str) -> List[str]:
    return re.split(r"(?<=[.!?])\s+(?=[A-Z])", text)


def _extract_tex_citations_from_file(
    tex_path: str, bib_lookup: Dict[str, dict]
) -> Tuple[str, List[dict]]:
    """Extract citation usages from a single .tex file, grouped by citation key."""
    with open(tex_path, "r", encoding="utf-8") as f:
        raw = f.read()

    # Remove comment lines
    cleaned_lines = [l for l in raw.split("\n") if not l.strip().startswith("%")]
    cleaned = " ".join(cleaned_lines)

    sentences = _split_sentences(cleaned)
    cite_pattern = re.compile(r"\\cite[a-z]*\{([^}]+)\}")
    citations_grouped: Dict[str, dict] = {}

    for i, sentence in enumerate(sentences):
        matches = cite_pattern.findall(sentence)
        if not matches:
            continue
        start_idx = max(0, i - 2)
        end_idx = min(len(sentences), i + 4)
        context = " ".join(sentences[start_idx:end_idx])

        # Find section heading
        snippet = sentence[:100]
        pos = raw.find(snippet)
        if pos == -1:
            pos = len(raw) // 2
        section = _find_section_heading(raw, pos)

        for match_str in matches:
            for key in (k.strip() for k in match_str.split(",")):
                if key not in bib_lookup:
                    continue
                if key not in citations_grouped:
                    citations_grouped[key] = {
                        "key": key,
                        "bib_entry": bib_lookup[key],
                        "ref_text": _format_bib_entry(bib_lookup[key]),
                        "usages": [],
                    }
                citations_grouped[key]["usages"].append({
                    "sentence": sentence,
                    "context": context,
                    "section": section,
                })

    # Build output
    citations = []
    for key, data in citations_grouped.items():
        unique_contexts = []
        seen = set()
        for usage in data["usages"]:
            ctx = usage["context"]
            section = usage["section"]
            if ctx not in seen:
                seen.add(ctx)
                unique_contexts.append(f"--- {section} ---\n{ctx}")

        citations.append({
            "key": key,
            "ref_text": data["ref_text"],
            "usage_count": len(data["usages"]),
            "unique_context_count": len(unique_contexts),
            "context_block": "\n\n".join(unique_contexts),
        })

    return cleaned, citations


def extract_from_tex(tex_paths: List[str], bib_path: str) -> List[dict]:
    """Extract citations from LaTeX (.tex) files with bibliography (.bib).

    Args:
        tex_paths: List of .tex file paths.
        bib_path: Path to the .bib bibliography file.

    Returns:
        List of citation dicts with keys: key, ref_text, usage_count,
        unique_context_count, context_block.
    """
    bib_lookup = _parse_bib(bib_path)
    all_citations: Dict[str, dict] = {}

    for tex_path in tex_paths:
        _, citations = _extract_tex_citations_from_file(tex_path, bib_lookup)
        for cit in citations:
            key = cit["key"]
            if key not in all_citations:
                all_citations[key] = cit
            else:
                # Merge usages from multiple .tex files
                all_citations[key]["usage_count"] += cit["usage_count"]
                all_citations[key]["unique_context_count"] += cit["unique_context_count"]
                all_citations[key]["context_block"] += "\n\n" + cit["context_block"]

    return list(all_citations.values())
